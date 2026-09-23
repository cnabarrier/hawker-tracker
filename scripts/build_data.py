"""Convert Zumi's hawker rating workbook (.xlsx/.xlsm) into the site's data.json shape.

build(workbook_bytes) -> (centres, warnings). Each centre is shaped exactly as index.html expects:
  unrated: {"n", "name", "addr"} plus "best"/"note"/"d" only if filled in
  rated:   {"n", "name", "addr", "s": [5 scores], "t", "rk", "best", "note"} plus "d" if filled in

Anything that looks wrong raises SheetError, so a bad sheet is never published.
"""
import datetime as dt
import io
import re

import openpyxl

SHEET = "Hawker Centre Ratings"

# Header prefixes (case-insensitive) -> field. Matching by prefix tolerates emoji/suffix tweaks.
COLUMNS = [
    ("s/n", "n"),
    ("hawker centre name", "name"),
    ("address", "addr"),
    ("food variation", "s0"),
    ("quality", "s1"),
    ("price", "s2"),
    ("comfy", "s3"),
    ("uniqueness", "s4"),
    ("total score", "t"),
    ("overall rank", "rk"),
    ("best time", "best"),
    ("notes", "note"),
    ("date visited", "d"),
]
SCORES = ["s0", "s1", "s2", "s3", "s4"]
LABELS = {"s0": "Food Variation", "s1": "Quality", "s2": "Price", "s3": "Comfy", "s4": "Uniqueness"}

# Same thresholds as her Rank formula (Summary!F23:G27). Below 21 her sheet counts it as F.
TIERS = [(45, "S"), (40, "A"), (35, "B"), (30, "C"), (21, "D")]

MONTHS = {}
for i, full in enumerate(["january", "february", "march", "april", "may", "june", "july",
                          "august", "september", "october", "november", "december"], 1):
    MONTHS[full] = MONTHS[full[:3]] = i
MONTHS["sept"] = 9

MIN_CENTRES = 100
EPS = 1e-9


class SheetError(Exception):
    pass


def tier(total):
    for floor, rk in TIERS:
        if total >= floor:
            return rk
    return "F"


def num(v):
    """Keep whole numbers as ints so the JSON reads like the sheet."""
    v = float(v)
    return int(v) if v.is_integer() else v


def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def clean_text(v):
    """Trim, and turn line breaks inside a cell into ' · ' (as the original snapshot did)."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\r\n", "\n").replace("\r", "\n").strip()
    return re.sub(r"\s*\n\s*", " · ", s)


def _date(y, m, d, where, raw):
    try:
        return dt.date(y, m, d)
    except ValueError:
        raise SheetError(f"{where}: {raw!r} is not a real date")


def parse_date(v, where, today):
    """ISO yyyy-mm-dd, or None if blank. Raises SheetError unless the date is unambiguous."""
    if blank(v):
        return None
    if isinstance(v, dt.datetime):
        d = v.date()
    elif isinstance(v, dt.date):
        d = v
    elif isinstance(v, (int, float)) and not isinstance(v, bool):
        if not 40000 < v < 60000:
            raise SheetError(f"{where}: Date Visited {v!r} is a number, not a date")
        d = (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    else:
        s = str(v).strip()
        if m := re.fullmatch(r"(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]+)\.?,?\s+(\d{4})", s):
            month = MONTHS.get(m[2].lower())
            if not month:
                raise SheetError(f"{where}: can't read the month in Date Visited {s!r}")
            d = _date(int(m[3]), month, int(m[1]), where, s)
        elif m := re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s):
            d = _date(int(m[1]), int(m[2]), int(m[3]), where, s)
        elif m := re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s):  # Singapore order: day/month/year
            d = _date(int(m[3]), int(m[2]), int(m[1]), where, s)
        else:
            raise SheetError(f"{where}: can't read Date Visited {s!r}")
    if not dt.date(2024, 1, 1) <= d <= today + dt.timedelta(days=1):
        raise SheetError(f"{where}: Date Visited {d.isoformat()} is outside the expected range")
    return d.isoformat()


def score(v, where, field):
    if blank(v):
        return None
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        raise SheetError(f"{where}: {LABELS[field]} score {v!r} is not a number")
    if not 0 <= v <= 10:
        raise SheetError(f"{where}: {LABELS[field]} score {v} is outside 0-10")
    return num(v)


def build(workbook_bytes, today=None):
    today = today or (dt.datetime.now(dt.timezone.utc) + dt.timedelta(hours=8)).date()  # Singapore
    try:
        wb = openpyxl.load_workbook(io.BytesIO(workbook_bytes), data_only=True)
    except Exception as e:
        raise SheetError(f"downloaded file is not a readable workbook ({e.__class__.__name__})")
    if SHEET not in wb.sheetnames:
        raise SheetError(f"tab {SHEET!r} not found (tabs: {wb.sheetnames})")
    rows = list(wb[SHEET].iter_rows(values_only=True))

    hdr_i = next((i for i, r in enumerate(rows[:15])
                  if r and isinstance(r[0], str) and r[0].strip().lower() == "s/n"), None)
    if hdr_i is None:
        raise SheetError("header row (starting with 'S/N') not found")
    headers = [str(h).strip().lower() if h is not None else "" for h in rows[hdr_i]]
    col = {}
    for prefix, field in COLUMNS:
        hits = [i for i, h in enumerate(headers) if h.startswith(prefix)]
        if len(hits) != 1:
            raise SheetError(f"expected exactly one column starting with {prefix!r}, found {len(hits)}")
        col[field] = hits[0]

    warnings, centres = [], []
    seen_n, seen_name = set(), set()
    body = rows[hdr_i + 1:]
    # Trailing empty rows are fine; an empty row followed by more data is not.
    last = max((i for i, r in enumerate(body) if any(not blank(v) for v in r)), default=-1)
    for i, r in enumerate(body[:last + 1]):
        rownum = hdr_i + 2 + i
        if all(blank(v) for v in r):
            raise SheetError(f"row {rownum} is empty but there are centres below it")
        get = lambda f: r[col[f]] if col[f] < len(r) else None
        name = clean_text(get("name"))
        where = f"row {rownum} ({name or 'no name'})"
        n = get("n")
        if isinstance(n, bool) or not isinstance(n, (int, float)) or not float(n).is_integer():
            raise SheetError(f"{where}: S/N {n!r} is not a whole number")
        n = int(n)
        if not name:
            raise SheetError(f"{where}: centre name is empty")
        if n in seen_n:
            raise SheetError(f"{where}: S/N {n} is used twice")
        if name.lower() in seen_name:
            raise SheetError(f"{where}: centre name appears twice")
        seen_n.add(n)
        seen_name.add(name.lower())

        c = {"n": n, "name": name, "addr": clean_text(get("addr"))}
        s = [score(get(f), where, f) for f in SCORES]
        best, note = clean_text(get("best")), clean_text(get("note"))
        d = parse_date(get("d"), where, today)

        if all(v is not None for v in s):
            total = sum(s)
            t_cell, rk_cell = get("t"), get("rk")
            if blank(t_cell):
                warnings.append(f"{where}: all 5 scores filled but Total is blank; using the sum {num(total)}")
            elif isinstance(t_cell, bool) or not isinstance(t_cell, (int, float)):
                raise SheetError(f"{where}: Total {t_cell!r} is not a number")
            elif abs(t_cell - total) > EPS:
                raise SheetError(f"{where}: Total {num(t_cell)} does not equal the sum of the 5 scores ({num(total)})")
            rk = tier(total)
            rk_txt = clean_text(rk_cell).upper()
            if rk_txt in ("", "0"):
                if rk != "F":  # her formula leaves <21 blank; anything else blank is unexpected
                    warnings.append(f"{where}: Overall Rank is blank; using {rk} from Total {num(total)}")
            elif rk_txt != rk:
                raise SheetError(f"{where}: Overall Rank {rk_txt!r} does not match Total {num(total)} (should be {rk})")
            c.update(s=s, t=num(total), rk=rk, best=best, note=note)
            if d:
                c["d"] = d
        else:
            if any(v is not None for v in s):
                warnings.append(f"{where}: only {sum(v is not None for v in s)}/5 scores filled; "
                                "kept as not rated until all 5 are in")
            for k, v in (("best", best), ("note", note), ("d", d)):
                if v:
                    c[k] = v
        centres.append(c)

    if len(centres) < MIN_CENTRES:
        raise SheetError(f"only {len(centres)} centres found (expected at least {MIN_CENTRES})")
    return centres, warnings
